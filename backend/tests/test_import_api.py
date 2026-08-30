"""Tests for the staging-only /api/import/spreadsheet endpoint.

Two concerns, tested separately:
1. The env gate — the real `app.main.app` (loaded with `ENABLE_SPREADSHEET_IMPORT=false`
   from backend/.env, as it is in every non-staging environment) must not expose
   this route at all.
2. The endpoint's actual behavior — exercised against a throwaway app instance
   with the router force-included and `get_db` overridden to an in-memory SQLite
   session, so this doesn't touch the real database either way.
"""

import io
import sqlite3
from datetime import date

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api import import_data
from app.db import get_db
from app.main import app as real_app
from app.models import Account, Base, TradeType


def test_import_route_is_absent_from_the_real_app_by_default():
    client = TestClient(real_app)
    response = client.post("/api/import/spreadsheet", params={"account": "rule1", "year": 2026})
    assert response.status_code == 404


@pytest.fixture
def staging_client():
    # FastAPI runs sync endpoints in a worker thread, so an in-memory SQLite
    # DB needs a single shared connection (StaticPool + check_same_thread=False)
    # rather than the usual one-connection-per-thread default.
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _register_now(dbapi_connection: sqlite3.Connection, _):
        dbapi_connection.create_function("now", 0, lambda: "2026-01-01 00:00:00")

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)

    with factory() as seed:
        seed.add(Account(account_name="Rule 1", account_type="INVESTMENT", start_date=date(2020, 1, 1)))
        seed.add(
            TradeType(
                type_name="ROCT PUT",
                category="OPTIONS",
                is_credit=True,
                requires_expiration=True,
                requires_strike=True,
                requires_contracts=True,
            )
        )
        seed.commit()

    test_app = FastAPI()
    test_app.include_router(import_data.router)

    def _override_get_db():
        db = factory()
        try:
            yield db
        finally:
            db.close()

    test_app.dependency_overrides[get_db] = _override_get_db
    return TestClient(test_app)


def _tiny_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "TRADES 2026"
    labels = {
        "TRADE DATE": 3,
        "UNDERLYING": 4,
        "PRICE": 13,
        "SHORT STRIKE": 23,
        "RESULT": 65,
    }
    for label, row in labels.items():
        ws.cell(row=row, column=2, value=label)

    ws.cell(row=1, column=5, value="SLV ROCT PUT")
    ws.cell(row=3, column=5, value=date(2026, 1, 2))
    ws.cell(row=4, column=5, value="SLV")
    ws.cell(row=13, column=5, value=85.04)
    ws.cell(row=23, column=5, value=55.5)
    ws.cell(row=65, column=5, value="EXPIRED")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_endpoint_loads_a_trade(staging_client):
    response = staging_client.post(
        "/api/import/spreadsheet",
        params={"account": "rule1", "year": 2026},
        files={"file": ("OKW_2026.xlsx", _tiny_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["trades_created"] == 1
    assert body["events_created"] == 2
    assert body["needs_review"] == 0


def test_import_endpoint_dry_run_does_not_write(staging_client):
    response = staging_client.post(
        "/api/import/spreadsheet",
        params={"account": "rule1", "year": 2026, "dry_run": True},
        files={"file": ("OKW_2026.xlsx", _tiny_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    assert response.json()["trades_created"] == 0

    # A real (non-dry-run) import should still see this as fresh, unimported data.
    response2 = staging_client.post(
        "/api/import/spreadsheet",
        params={"account": "rule1", "year": 2026},
        files={"file": ("OKW_2026.xlsx", _tiny_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response2.status_code == 200
    assert response2.json()["trades_created"] == 1


def test_import_endpoint_rejects_unknown_account(staging_client):
    response = staging_client.post(
        "/api/import/spreadsheet",
        params={"account": "not_a_real_account", "year": 2026},
        files={"file": ("OKW_2026.xlsx", _tiny_workbook_bytes(), "application/octet-stream")},
    )
    assert response.status_code == 400
