"""Pure parser for the OKW-style trade-tracking workbook.

Layout discovered from the real `OKW_2026.xlsx` (`TRADES 2026` / `ROTH TRADES
2026` sheets): each trade occupies one *value column* (not a row) — column
B holds the row's field label for the whole sheet, and each trade's actual
values live in whatever column its row-1 header cell is in (e.g. `E`, `G`,
`I`, ... — spacing between trades is **not** reliably 2 columns apart, so
block columns are discovered by scanning row 1 for non-empty cells, not
assumed from a fixed stride).

Deliberately has zero DB/SQLAlchemy dependency — this module only turns a
worksheet into `ParsedTrade` records. `scripts/import_historical.py` (CLI)
and `app/api/import_data.py` (staging endpoint) both call this, then handle
persistence/idempotency separately. Keeping this pure makes it cheap to unit
test against a small synthetic workbook instead of the real, private
spreadsheet file.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

import re

from openpyxl.worksheet.worksheet import Worksheet

# Row-3..90 field labels, as they appear literally in column B. Matched
# case-insensitively with whitespace collapsed, since the same label has
# been observed with inconsistent leading/trailing spaces across sheets
# (e.g. `"EXPIRATION DATE "` vs `"EXPIRATION DATE"`).
LABELS = {
    "trade_date": "TRADE DATE",
    "underlying": "UNDERLYING",
    "price": "PRICE",
    "dte": "DAYS TO EXPIRATION (DTE)",
    "expiration_date": "EXPIRATION DATE",
    "short_strike": "SHORT STRIKE",
    "long_strike": "LONG STRIKE",
    "credit_debit": "CREDIT/(DEBIT)",
    "commission_per_share": "COMMISSION (PER SHARE)",
    "net_credit_per_share": "NET CREDIT (NC)",
    "risk_capital_per_share": "RISK CAPITAL (RC)",
    "margin_percent": "MARGIN %",
    "margin_capital": "MARGIN CAPITAL",
    "arorc": "ANNUALIZED RORC (ARORC)",
    "probability_of_winning": "PROBABILITY OF WINNING",
    "delta": "DELTA",
    "contracts": "ACTUAL CONTRACTS",
    "shares": "SHARES",
    "result": "RESULT",
    "result_date": "RESULT DATE (CLOSED OR EXPIRED)",
    "closing_debit": "CLOSING DEBIT",
    "total_debit": "TOTAL DEBIT",
    "result_net_credit": "NET CREDIT/(DEBIT)",
    "final_arorc": "FINAL ARORC",
    "notes": "NOTES",
}

# Extra column-B spellings that map onto LABELS keys (UI rename / older files).
LABEL_ALIASES = {
    "PROB OTM": "probability_of_winning",
    "PROBABILITY OF WINNING (PROB OTM)": "probability_of_winning",
}

# RESULT cell text -> trade_events.event_type. Anything not in this map
# (including blank/None, meaning still open) is left unmapped — the caller
# decides what an unrecognized non-blank RESULT means (flag for review,
# never guess).
RESULT_TO_EVENT_TYPE = {
    "EXPIRED": "EXPIRE",
    "ASSIGNED": "ASSIGN",
    "CLOSED": "CLOSE",
    "ROLLED": "ROLL",
    "ROLL": "ROLL",
}

# RESULT values that explicitly mean "still open" — recognized, but
# intentionally produce no closing event_type and no review flag.
OPEN_RESULT_VALUES = {"OPEN", ""}

FIRST_DATA_COLUMN = 5  # column E — trade blocks never start before this


def _normalize_label(text: str) -> str:
    return " ".join(text.split()).upper()


@dataclass
class ParsedTrade:
    """One column-block's worth of data, before any DB lookups.

    `ticker` and `trade_type_name` are raw strings — matching them to
    `tickers`/`trade_types` rows (get-or-create / lookup) is the loader's
    job, not the parser's, since that requires a DB session.
    """

    sheet_name: str
    column_letter: str
    ticker: str
    trade_type_name: str
    trade_date: date | None
    price_per_share: Decimal | None
    days_to_expiration: int | None
    expiration_date: date | None
    short_strike: Decimal | None
    long_strike: Decimal | None
    credit_debit: Decimal | None
    commission_per_share: Decimal | None
    net_credit_per_share: Decimal | None
    risk_capital_per_share: Decimal | None
    margin_percent: Decimal | None
    margin_capital: Decimal | None
    arorc: Decimal | None
    delta: Decimal | None
    probability_of_winning: Decimal | None
    num_of_contracts: int | None
    num_of_shares: int | None
    result_raw: str | None
    result_date: date | None
    closing_debit: Decimal | None
    total_debit: Decimal | None
    result_net_credit: Decimal | None
    final_arorc: Decimal | None
    notes: str | None
    warnings: list[str] = field(default_factory=list)

    @property
    def event_type(self) -> str | None:
        """`trade_events.event_type` for the closing event, or `None` if
        the trade is still open (blank RESULT) or the RESULT text isn't
        recognized (see `warnings` in that case — never guess silently).
        """
        if not self.result_raw:
            return None
        return RESULT_TO_EVENT_TYPE.get(self.result_raw.strip().upper())

    @property
    def needs_review(self) -> bool:
        return bool(self.warnings)


def _to_decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, str) and value.strip().startswith("="):
        return None
    if isinstance(value, str) and value.strip().endswith("%"):
        try:
            return Decimal(value.strip()[:-1].strip()) / Decimal("100")
        except (InvalidOperation, ValueError):
            return None
    if isinstance(value, str) and not value.strip().replace(".", "").replace(
        "-", ""
    ).isdigit():
        # Non-numeric placeholders like "COVERED" (margin fields on covered
        # calls) or "-" (probability not yet computed) — not an error, just
        # not a number to carry into a NUMERIC column.
        return None
    try:
        if isinstance(value, float):
            return Decimal(str(round(value, 12)))
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _quantize(value: Decimal | None, quantum: Decimal) -> Decimal | None:
    if value is None:
        return None
    return value.quantize(quantum, rounding=ROUND_HALF_UP)


_HEADER_TYPE_RE = re.compile(r'"\s*([^"]+)"\s*$')


def _header_ticker_and_type(header_value: object, underlying: str) -> tuple[str, str]:
    """Row-1 is often a formula (`=E4&" ROCT PUT"`). Prefer the cached
    `SLV ROCT PUT` text; if we only have the formula, take the ticker from
    UNDERLYING and the type from the quoted suffix.
    """
    header_text = str(header_value).strip() if header_value not in (None, "") else ""
    if header_text.startswith("="):
        ticker = underlying.split()[0].strip() if underlying else ""
        match = _HEADER_TYPE_RE.search(header_text)
        trade_type_name = match.group(1).strip() if match else ""
        return ticker, trade_type_name
    parts = header_text.split(" ", 1)
    ticker = parts[0].strip()
    trade_type_name = parts[1].strip() if len(parts) > 1 else ""
    return ticker, trade_type_name


def _prob_otm_from_delta(delta: Decimal | None) -> Decimal | None:
    """OKW `PROBABILITY OF WINNING` is `IF(delta=0,"-",1-delta)` (Prob OTM)."""
    if delta is None or delta == 0:
        return None
    return (Decimal("1") - delta).quantize(Decimal("0.000001"))


def _to_int(value: object) -> int | None:
    dec = _to_decimal(value)
    return int(dec) if dec is not None else None


def _to_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if hasattr(value, "date"):  # datetime -> date
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _build_label_row_map(ws: Worksheet) -> dict[str, list[int]]:
    """Scan column B for each row's label, keyed by our normalized field
    names (`LABELS` above). DELTA / PROBABILITY OF WINNING appear twice
    (entry block and FINAL TRADE REVIEW); callers take the first numeric.
    """
    wanted = {_normalize_label(v): k for k, v in LABELS.items()}
    wanted.update({_normalize_label(alias): key for alias, key in LABEL_ALIASES.items()})
    found: dict[str, list[int]] = {}
    for row_num, row_cells in enumerate(
        ws.iter_rows(min_col=2, max_col=2, max_row=200), start=1
    ):
        value = row_cells[0].value
        if value is None or not isinstance(value, str):
            continue
        key = wanted.get(_normalize_label(value))
        if key:
            found.setdefault(key, []).append(row_num)
    return found


def _find_block_columns(ws: Worksheet) -> list[int]:
    """Every column with a non-empty row-1 cell at or after column E is the
    start of one trade's data. Spacing between blocks is irregular in the
    real workbook, so this can't assume a fixed stride.
    """
    columns = []
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    for col_num, cell in enumerate(header_row, start=1):
        if col_num >= FIRST_DATA_COLUMN and cell.value not in (None, ""):
            columns.append(col_num)
    return columns


def parse_trade_sheet(ws: Worksheet) -> list[ParsedTrade]:
    """Parse every trade block in one sheet (e.g. `TRADES 2026`)."""
    label_rows = _build_label_row_map(ws)
    missing = [k for k in ("trade_date", "underlying", "price") if k not in label_rows]
    if missing:
        raise ValueError(
            f"Sheet '{ws.title}' is missing expected row labels {missing} in column "
            "B — this doesn't look like an OKW trade sheet, or its layout has "
            "changed enough that LABELS in okw_parser.py needs updating."
        )

    def get(col: int, key: str) -> object:
        rows = label_rows.get(key) or []
        if not rows:
            return None
        return ws.cell(row=rows[0], column=col).value

    def get_numeric(col: int, key: str) -> Decimal | None:
        for row in label_rows.get(key) or []:
            parsed = _to_decimal(ws.cell(row=row, column=col).value)
            if parsed is not None:
                return parsed
        return None

    trades: list[ParsedTrade] = []
    for col in _find_block_columns(ws):
        header_cell = ws.cell(row=1, column=col)
        underlying = str(get(col, "underlying") or "").strip()
        ticker, trade_type_name = _header_ticker_and_type(header_cell.value, underlying)

        warnings: list[str] = []
        if not ticker:
            warnings.append(f"No header text found for column {header_cell.coordinate}")
        if (
            underlying
            and ticker
            and not str(header_cell.value or "").strip().startswith("=")
            and not underlying.upper().startswith(ticker.upper())
        ):
            warnings.append(
                f"UNDERLYING ({underlying!r}) doesn't match header ticker ({ticker!r})"
            )
        result_raw = get(col, "result")
        result_text = str(result_raw).strip() if result_raw not in (None, "") else None
        if (
            result_text
            and result_text.upper() not in RESULT_TO_EVENT_TYPE
            and result_text.upper() not in OPEN_RESULT_VALUES
        ):
            warnings.append(f"Unrecognized RESULT value: {result_text!r}")

        trades.append(
            ParsedTrade(
                sheet_name=ws.title,
                column_letter=header_cell.column_letter,
                ticker=ticker,
                trade_type_name=trade_type_name,
                trade_date=_to_date(get(col, "trade_date")),
                price_per_share=_to_decimal(get(col, "price")),
                days_to_expiration=_to_int(get(col, "dte")),
                expiration_date=_to_date(get(col, "expiration_date")),
                short_strike=_to_decimal(get(col, "short_strike")),
                long_strike=(
                    _to_decimal(get(col, "long_strike"))
                    if "SPREAD" in trade_type_name.upper()
                    else None
                ),
                credit_debit=_to_decimal(get(col, "credit_debit")),
                commission_per_share=_to_decimal(get(col, "commission_per_share")),
                net_credit_per_share=_to_decimal(get(col, "net_credit_per_share")),
                risk_capital_per_share=_to_decimal(get(col, "risk_capital_per_share")),
                margin_percent=_to_decimal(get(col, "margin_percent")),
                margin_capital=_to_decimal(get(col, "margin_capital")),
                arorc=_to_decimal(get(col, "arorc")),
                delta=get_numeric(col, "delta"),
                probability_of_winning=(
                    get_numeric(col, "probability_of_winning")
                    or _prob_otm_from_delta(get_numeric(col, "delta"))
                ),
                num_of_contracts=_to_int(get(col, "contracts")),
                num_of_shares=_to_int(get(col, "shares")),
                result_raw=result_text,
                result_date=_to_date(get(col, "result_date")),
                closing_debit=_to_decimal(get(col, "closing_debit")),
                total_debit=_to_decimal(get(col, "total_debit")),
                result_net_credit=_quantize(
                    get_numeric(col, "result_net_credit"), Decimal("0.01")
                ),
                final_arorc=_quantize(
                    get_numeric(col, "final_arorc"), Decimal("0.000001")
                ),
                notes=(str(n) if (n := get(col, "notes")) not in (None, "") else None),
                warnings=warnings,
            )
        )
    return trades
